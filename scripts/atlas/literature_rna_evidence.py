"""Curate Vasconcelos and Almeida publication evidence for the RNA atlas.

These sources are deliberately kept outside the bulk expression contract.  The
published identifiers and statistics are preserved, while a current atlas gene
is assigned only through the same exact/unique-stable-ID rules used by the
other external RNA layers.
"""

from __future__ import annotations

import hashlib
import re
from collections import defaultdict
from pathlib import Path

from external_rna_evidence import resolve


VASCONCELOS_URL = "https://doi.org/10.1038/s41598-017-10853-6"
ALMEIDA_URL = "https://doi.org/10.1016/j.exppara.2011.06.010"
REFERENCE = "Schistosoma_mansoni_SM_V10_WBPS19"
SOURCE_ASSEMBLY = "Schistosoma_mansoni_v5.2"

LNCRNA_FIELDS = [
    "source_study", "source_publication", "source_location", "source_record_id",
    "original_id", "lncrna_class", "source_assembly", "source_chromosome",
    "source_start", "source_end", "strand", "block_count", "block_sizes",
    "block_starts", "robust_candidate", "stage_differential_expression",
    "log2fc_somula_3h", "log2fc_somula_24h", "log2fc_adult_male",
    "log2fc_adult_female", "log2_cpm", "p_value", "fdr",
    "associated_original_gene_id", "associated_current_gene_id",
    "association_mapping_method", "association_mapping_status",
    "mapping_source", "mapping_source_version", "catalog_source_checksum",
    "stage_de_source_checksum", "robust_source_checksum", "curation_status", "notes",
]

LEGACY_FIELDS = [
    "source_study", "source_publication", "source_location", "source_record_id",
    "evidence_type", "original_id", "current_gene_id", "mapping_method",
    "mapping_source", "mapping_source_version", "mapping_status", "life_stage",
    "sex", "est_count", "source_contig_size_bp", "source_scaffold",
    "published_annotation", "assertion", "source_checksum", "curation_status", "notes",
]


def checksum(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def previous_id_index(atlas_genes: set[str], annotations: dict[str, dict[str, str]]) -> dict[str, set[str]]:
    previous: dict[str, set[str]] = defaultdict(set)
    for gene, annotation in annotations.items():
        if gene not in atlas_genes:
            continue
        for token in re.split(r"[,;]", annotation.get("previous_stable_id", "")):
            if token.strip():
                previous[token.strip()].add(gene)
    return previous


def xlsx_rows(path: Path, sheet_name: str, min_row: int):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        for excel_row, values in enumerate(sheet.iter_rows(min_row=min_row, values_only=True), min_row):
            yield excel_row, values
    finally:
        workbook.close()


def xls_rows(path: Path, min_row: int):
    import xlrd

    workbook = xlrd.open_workbook(path, on_demand=True)
    try:
        sheet = workbook.sheet_by_index(0)
        for zero_based_row in range(min_row - 1, sheet.nrows):
            yield zero_based_row + 1, sheet.row_values(zero_based_row)
    finally:
        workbook.release_resources()


def curate_vasconcelos(source_dir: Path, atlas_genes: set[str], annotations: dict[str, dict[str, str]]):
    catalog_path = source_dir / "vasconcelos_table_s1.xlsx"
    stage_path = source_dir / "vasconcelos_table_s4.xlsx"
    robust_path = source_dir / "vasconcelos_file_s2.xlsx"
    for path in (catalog_path, stage_path, robust_path):
        if not path.is_file():
            raise FileNotFoundError(f"Vasconcelos source table missing: {path}")

    source_hashes = {path: checksum(path) for path in (catalog_path, stage_path, robust_path)}
    robust = {
        clean(values[0])
        for _, values in xlsx_rows(robust_path, "File S2", 2)
        if values and clean(values[0]).startswith("SmLINC")
    }
    stage_de: dict[str, tuple[str, ...]] = {}
    for _, values in xlsx_rows(stage_path, "Plan1", 3):
        identifier = clean(values[0]) if values else ""
        if identifier.startswith("SmLINC"):
            stage_de[identifier] = tuple(clean(value) for value in values[1:8])

    antisense_links: dict[str, list[str]] = defaultdict(list)
    for _, values in xlsx_rows(catalog_path, "SmAS-Smp", 2):
        if len(values) < 2:
            continue
        lncrna, gene = clean(values[0]), clean(values[1])
        if lncrna and gene and gene not in antisense_links[lncrna]:
            antisense_links[lncrna].append(gene)

    previous = previous_id_index(atlas_genes, annotations)
    records: list[dict[str, str]] = []

    def add_catalog(sheet_name: str, lncrna_class: str):
        for excel_row, values in xlsx_rows(catalog_path, sheet_name, 3):
            if len(values) < 8:
                continue
            identifier = clean(values[3])
            if not identifier:
                continue
            original_genes = antisense_links.get(identifier, [])
            mapped_genes: list[str] = []
            methods: list[str] = []
            statuses: list[str] = []
            for original_gene in original_genes:
                gene, method, status = resolve(original_gene, atlas_genes, previous)
                methods.append(f"{original_gene}:{method}")
                statuses.append(status)
                if gene and gene not in mapped_genes:
                    mapped_genes.append(gene)
            if not original_genes:
                association_status = "NOT_APPLICABLE"
            elif statuses and all(status == "MAPPED" for status in statuses):
                association_status = "MAPPED"
            elif "MAPPED" in statuses:
                association_status = "PARTIAL"
            elif "AMBIGUOUS" in statuses:
                association_status = "AMBIGUOUS"
            else:
                association_status = "UNMAPPED"
            de = stage_de.get(identifier, ("", "", "", "", "", "", ""))
            records.append({
                "source_study": "VASCONCELOS_2017_LNCRNA_CATALOG",
                "source_publication": VASCONCELOS_URL,
                "source_location": f"Supplementary Table S1; sheet {sheet_name}",
                "source_record_id": f"S1:{sheet_name}:row:{excel_row}",
                "original_id": identifier,
                "lncrna_class": lncrna_class,
                "source_assembly": SOURCE_ASSEMBLY,
                "source_chromosome": clean(values[0]),
                "source_start": clean(values[1]),
                "source_end": clean(values[2]),
                "strand": clean(values[4]),
                "block_count": clean(values[5]),
                "block_sizes": clean(values[6]),
                "block_starts": clean(values[7]),
                "robust_candidate": "YES" if identifier in robust else "NO",
                "stage_differential_expression": "YES" if identifier in stage_de else "NO",
                "log2fc_somula_3h": de[0],
                "log2fc_somula_24h": de[1],
                "log2fc_adult_male": de[2],
                "log2fc_adult_female": de[3],
                "log2_cpm": de[4],
                "p_value": de[5],
                "fdr": de[6],
                "associated_original_gene_id": ";".join(original_genes),
                "associated_current_gene_id": ";".join(mapped_genes),
                "association_mapping_method": ";".join(methods) if methods else "NOT_APPLICABLE_DISTINCT_LNCRNA_ENTITY",
                "association_mapping_status": association_status,
                "mapping_source": "WBPS19 gene annotations" if original_genes else "",
                "mapping_source_version": REFERENCE if original_genes else "",
                "catalog_source_checksum": source_hashes[catalog_path],
                "stage_de_source_checksum": source_hashes[stage_path] if identifier in stage_de else "",
                "robust_source_checksum": source_hashes[robust_path] if identifier in robust else "",
                "curation_status": "PUBLISHED_SOURCE_COORDINATES_NOT_LIFTED_TO_SM_V10",
                "notes": (
                    "Published multi-exon lncRNA model on the v5.2 assembly. It is an annotation/evidence feature, "
                    "not a WBPS19 gene, atlas sample, TPM measurement, or HelixForge differential-expression result."
                ),
            })

    add_catalog("SmLINC", "LINC_RNA")
    add_catalog("SmAS", "ANTISENSE_LNCRNA")
    return records, [catalog_path, stage_path, robust_path]


def curate_almeida(source_dir: Path, atlas_genes: set[str], annotations: dict[str, dict[str, str]]):
    confirmed_path = source_dir / "almeida_supplementary_table_2.xls"
    first_path = source_dir / "almeida_supplementary_table_3.xls"
    contig_path = source_dir / "almeida_supplementary_table_4.xls"
    for path in (confirmed_path, first_path, contig_path):
        if not path.is_file():
            raise FileNotFoundError(f"Almeida source table missing: {path}")
    source_hashes = {path: checksum(path) for path in (confirmed_path, first_path, contig_path)}
    previous = previous_id_index(atlas_genes, annotations)
    records: list[dict[str, str]] = []

    def add_gene_table(path: Path, table_name: str, evidence_type: str, assertion: str):
        for excel_row, values in xls_rows(path, 4):
            if len(values) < 4:
                continue
            published_id = clean(values[1])
            if not published_id.startswith("Smp_"):
                continue
            original_gene = re.split(r":mRNA", published_id, maxsplit=1)[0]
            gene, method, status = resolve(original_gene, atlas_genes, previous)
            records.append({
                "source_study": "PRJNA64591",
                "source_publication": ALMEIDA_URL,
                "source_location": table_name,
                "source_record_id": f"{table_name.replace(' ', '_')}:row:{excel_row}",
                "evidence_type": evidence_type,
                "original_id": published_id,
                "current_gene_id": gene,
                "mapping_method": method,
                "mapping_source": "WBPS19 gene annotations",
                "mapping_source_version": REFERENCE,
                "mapping_status": status,
                "life_stage": "adult",
                "sex": "male",
                "est_count": clean(values[2]),
                "source_contig_size_bp": "",
                "source_scaffold": "",
                "published_annotation": clean(values[3]),
                "assertion": assertion,
                "source_checksum": source_hashes[path],
                "curation_status": "PUBLISHED_CLAIM_NOT_REANALYZED",
                "notes": "Roche 454 directional transcript evidence; published read counts are not atlas TPM values.",
            })

    add_gene_table(
        confirmed_path,
        "Supplementary Table 2",
        "PUBLISHED_ADULT_MALE_TRANSCRIPTION_SUPPORT",
        "Published adult-male EST support for a v5-era gene prediction.",
    )
    add_gene_table(
        first_path,
        "Supplementary Table 3",
        "PUBLISHED_FIRST_ADULT_MALE_TRANSCRIPTION_SUPPORT",
        "Published first transcript evidence for a v5-era prediction in adult male worms.",
    )

    for excel_row, values in xls_rows(contig_path, 4):
        if len(values) < 4:
            continue
        contig = clean(values[1])
        if not contig:
            continue
        scaffold_match = re.search(r"Smp_scaff\d+", contig)
        records.append({
            "source_study": "PRJNA64591",
            "source_publication": ALMEIDA_URL,
            "source_location": "Supplementary Table 4",
            "source_record_id": f"Supplementary_Table_4:row:{excel_row}",
            "evidence_type": "PUBLISHED_NOVEL_TRANSCRIBED_CONTIG",
            "original_id": contig,
            "current_gene_id": "",
            "mapping_method": "NO_STABLE_GENE_ID_SEQUENCE_REMAPPING_REQUIRED",
            "mapping_source": "",
            "mapping_source_version": "",
            "mapping_status": "UNMAPPED",
            "life_stage": "adult",
            "sex": "male",
            "est_count": clean(values[3]),
            "source_contig_size_bp": clean(values[2]),
            "source_scaffold": scaffold_match.group(0) if scaffold_match else "",
            "published_annotation": clean(values[6]) if len(values) > 6 else "",
            "assertion": "Published transcribed contig in a v5-era region without a predicted gene.",
            "source_checksum": source_hashes[contig_path],
            "curation_status": "SOURCE_CONTIG_NOT_REASSIGNED_TO_SM_V10",
            "notes": (
                "No current gene is assigned from the historical scaffold name. Sequence remapping to SM_V10 "
                "and a reviewed locus-overlap rule are required before gene-level attachment."
            ),
        })
    return records, [confirmed_path, first_path, contig_path]


def curate(source_dir: Path, atlas_genes: set[str], annotations: dict[str, dict[str, str]]):
    lncrna_rows, vasconcelos_sources = curate_vasconcelos(source_dir, atlas_genes, annotations)
    legacy_rows, almeida_sources = curate_almeida(source_dir, atlas_genes, annotations)
    return lncrna_rows, legacy_rows, vasconcelos_sources + almeida_sources

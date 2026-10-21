"""Curate published RNA evidence without treating it as bulk atlas data."""

from __future__ import annotations

import hashlib
import re
from collections import defaultdict
from pathlib import Path


ANDERSON_URL = "https://journals.plos.org/plosntds/article?id=10.1371/journal.pntd.0004334"
WANG_URL = "https://elifesciences.org/articles/35449"
REFERENCE = "Schistosoma_mansoni_SM_V10_WBPS19"
ANNOTATION_FIELDS = [
    "source_study", "source_publication", "source_location", "source_record_id",
    "evidence_type", "original_id", "current_gene_id", "mapping_method",
    "mapping_source", "mapping_source_version", "mapping_status", "stage_or_context",
    "assertion", "source_checksum", "curation_status", "notes",
]
MARKER_FIELDS = [
    "source_study", "source_publication", "source_location", "source_record_id",
    "original_id", "current_gene_id", "mapping_method", "mapping_source",
    "mapping_source_version", "mapping_status", "cell_population",
    "lineage_or_state", "life_stage", "marker_class", "reported_statistic",
    "assertion", "source_checksum", "curation_status", "notes",
]
def checksum(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def rows(path: Path, header_row: int):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        for _ in range(header_row - 1):
            next(iterator)
        fields = [str(value).strip() if value is not None else "" for value in next(iterator)]
        for excel_row, values in enumerate(iterator, header_row + 1):
            if values[0] is None:
                continue
            yield excel_row, dict(zip(fields, values))
    finally:
        workbook.close()


def clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def resolve(identifier: str, atlas_genes: set[str], previous: dict[str, set[str]]) -> tuple[str, str, str]:
    original = identifier.strip()
    normalized = re.sub(r"(?i)^smp_", "Smp_", original)
    candidate = re.sub(r"\.\d+$", "", normalized)
    if normalized in atlas_genes:
        return normalized, "EXACT_ATLAS_ID", "MAPPED"
    if candidate != normalized and candidate in atlas_genes:
        return candidate, "VERSION_SUFFIX_REMOVED_EXACT_ATLAS_ID", "MAPPED"
    matches = previous.get(normalized, set()) | previous.get(candidate, set())
    if len(matches) == 1:
        return next(iter(matches)), "UNIQUE_WBPS19_PREVIOUS_STABLE_ID", "MAPPED"
    if len(matches) > 1:
        return "", "AMBIGUOUS_WBPS19_PREVIOUS_STABLE_ID", "AMBIGUOUS"
    return "", "NO_UNIQUE_WBPS19_ATLAS_MATCH", "UNMAPPED"


def curate(source_dir: Path, atlas_genes: set[str], annotations: dict[str, dict[str, str]]):
    previous: dict[str, set[str]] = defaultdict(set)
    for gene, annotation in annotations.items():
        if gene not in atlas_genes:
            continue
        for token in re.split(r"[,;]", annotation.get("previous_stable_id", "")):
            if token.strip():
                previous[token.strip()].add(gene)

    source_names = ["anderson_s2", "anderson_s4", "anderson_s5", "anderson_s6", "wang_s1", "wang_s3"]
    paths = {name: source_dir / f"{name}.xlsx" for name in source_names}
    for path in paths.values():
        if not path.is_file():
            raise FileNotFoundError(f"published supplementary table missing: {path}")
    hashes = {name: checksum(path) for name, path in paths.items()}
    annotation_rows: list[dict[str, str]] = []
    marker_rows: list[dict[str, str]] = []

    def add_annotation(table: str, excel_row: int, identifier: str, kind: str, context: str, assertion: str, notes: str):
        gene, method, status = resolve(identifier, atlas_genes, previous)
        annotation_rows.append({
            "source_study": "PRJNA294789", "source_publication": ANDERSON_URL,
            "source_location": f"Supporting Information {table.upper()}",
            "source_record_id": f"{table.upper()}:row:{excel_row}", "evidence_type": kind,
            "original_id": identifier, "current_gene_id": gene, "mapping_method": method,
            "mapping_source": "WBPS19 gene annotations", "mapping_source_version": REFERENCE,
            "mapping_status": status, "stage_or_context": context,
            "assertion": assertion, "source_checksum": hashes[f"anderson_{table}"],
            "curation_status": "PUBLISHED_CLAIM_NOT_REANALYZED", "notes": notes,
        })

    for number, row in rows(paths["anderson_s2"], 2):
        flag = clean(row.get("Most highly expressed flag"))
        if flag in {"Egg_High", "Female_High", "Male_High"}:
            add_annotation("s2", number, clean(row["Smp_gene"]), "PUBLISHED_STAGE_ENRICHMENT",
                           flag.removesuffix("_High").lower(),
                           f"Published {flag} label; not an atlas DESeq2 result.",
                           f"RNA-seq contig: {clean(row.get('Contig RNA-Seq'))}; original Roche 454/NOISeq analysis.")
    for number, row in rows(paths["anderson_s4"], 2):
        extension_5 = clean(row.get("Extended_5'UTR")).upper() == "YES"
        extension_3 = clean(row.get("Extended_3'UTR")).upper() == "YES"
        if extension_5 or extension_3:
            ends = ", ".join(end for end, present in (("5'", extension_5), ("3'", extension_3)) if present)
            add_annotation("s4", number, clean(row["Smp_gene"]), "PUBLISHED_GENE_MODEL_EXTENSION",
                           "transcript boundary", f"Published {ends} gene-model extension claim.",
                           f"Contig: {clean(row.get('Contig_name'))}; published coordinates: {clean(row.get('Contig genomic coordinate'))}:{clean(row.get('Start_contig'))}-{clean(row.get('End_contig'))}.")
    for number, row in rows(paths["anderson_s5"], 2):
        add_annotation("s5", number, clean(row["Smp_gene"]), "PUBLISHED_MODEL_RESTRUCTURING",
                       "gene model", "Published restructured gene-model association; merge versus split is not inferred here.",
                       f"Contig: {clean(row.get('Contig_name'))}; published coordinates: {clean(row.get('contig genomic coordinate'))}:{clean(row.get('Start_contig'))}-{clean(row.get('End_contig'))}.")
    for number, row in rows(paths["anderson_s6"], 2):
        add_annotation("s6", number, clean(row["Contig_Name"]), "PUTATIVE_NOVEL_GENE_CLAIM",
                       "putative novel transcript", "Published putative novel gene; no WBPS19 gene assignment without unique mapping.",
                       f"CDD: {clean(row.get('CDD_Accession'))}; Pfam: {clean(row.get('Pfam_ID'))}; published coordinates: {clean(row.get('contig genomic coordinate'))}:{clean(row.get('Start_contig'))}-{clean(row.get('End_contig'))}.")

    published_names = {clean(row.get("Feature ID")): clean(row.get("Gene name")) for _, row in rows(paths["wang_s3"], 1)}
    for number, row in rows(paths["wang_s1"], 1):
        identifier = clean(row["Feature ID"])
        gene, method, status = resolve(identifier, atlas_genes, previous)
        statistic = "; ".join(f"{cell} TPM={clean(row.get(f'{cell} (TPM)'))}" for cell in ("φ", "δ", "κ"))
        marker_rows.append({
            "source_study": "PRJNA395457", "source_publication": WANG_URL,
            "source_location": "Supplementary file 1; gene names checked against supplementary file 3",
            "source_record_id": f"S1:row:{number}", "original_id": identifier,
            "current_gene_id": gene, "mapping_method": method,
            "mapping_source": "WBPS19 gene annotations", "mapping_source_version": REFERENCE,
            "mapping_status": status, "cell_population": "sporocyst stem cells",
            "lineage_or_state": "φ / δ / κ classes", "life_stage": "sporocyst",
            "marker_class": "PUBLISHED_CLASS_DEPENDENT_TRANSCRIPT", "reported_statistic": statistic,
            "assertion": "Published class-dependent transcript (>4-fold, p<0.01) among sporocyst stem-cell classes; not a bulk differential-expression result.",
            "source_checksum": hashes["wang_s1"], "curation_status": "PUBLISHED_CLAIM_NOT_REANALYZED",
            "notes": f"Genome v5 annotation: {clean(row.get('Annotation (S. mansoni genome v.5)'))}; published gene name: {published_names.get(identifier) or clean(row.get('Gene name'))}; S3 checksum: {hashes['wang_s3']}.",
        })
    return annotation_rows, marker_rows, list(paths.values())
